"""CopyChief quick mode: one-off jobs with no project folder.

Used by the copychief-quick skill. Everything happens in one scratch folder (the
chat's workspace); nothing is filed in a client or projects folder. The shared
engines (build_review_tracked.py, docx_lib.py) do the rendering, so the house
formatting stays identical to the folder workflows.

Commands
  python tools/quick.py extract <input> <work_dir>
      Turn a .docx / .md / .txt / .html / .pdf draft into numbered paragraphs.
      Writes <work_dir>/source.json and prints "[n] (style) text" for each one.

  python tools/quick.py tracked <work_dir> <out.docx>
      Read <work_dir>/source.json + <work_dir>/edits.json, build the tracked-changes
      review (.docx + .html redline), then verify it (see below).

  python tools/quick.py doc <blocks.json> <out.docx>
      Render a prep, scorecard or draft (docx_lib block list + comments) to a
      brand-formatted .docx and a matching .md.

  python tools/quick.py b64 <file>
      Print the file as base64 (for the Google Drive connector's upload call).

edits.json (the editor writes ONLY what changes; unchanged paragraphs are left out)
  {
    "title": "Article title",
    "comments": {"1": ["Short label", "Full reasoning"], ...},
    "edits": [
      {"n": 4, "r": "Revised paragraph text", "c": [1]},    # rewrite paragraph 4
      {"n": 7, "delete": true, "c": [2]},                  # delete paragraph 7
      {"after": 7, "r": "New paragraph", "style": "p", "c": [3]},  # insert after 7 (0 = top)
      {"n": 9, "c": [4]},                                  # comment only, no text change
      {"n": 2, "style": "h2", "r": "New heading"}          # style may change too
    ]
  }
  Text may use **bold** and [links](https://...). Styles: h1 h2 h3 p li.
  The script diffs original vs revised word by word, so reject-all always restores
  the author's exact wording and accept-all leaves clean copy. It checks both.

Why a diff instead of hand-written tracked tokens: the editor writes plain revised
text (far fewer tokens, no chance of mangling the original), and correctness of
the tracked changes is guaranteed by code instead of by care.
"""

from __future__ import annotations

import base64
import difflib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+|mailto:[^)\s]+)\)")
UNIT_RE = re.compile(r"\[[^\]]+\]\((?:https?://|mailto:)[^)\s]+\)|\*\*.+?\*\*|\s+|[^\s]+")


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------

def plain(text: str) -> str:
    """Markup-free text: links -> their text, bold markers removed, spaces collapsed."""
    text = LINK_RE.sub(r"\1", text)
    text = text.replace("**", "")
    return re.sub(r"\s+", " ", text).strip()


def units(text: str) -> list[tuple]:
    """Split text into diff units: (kind, text, url). kind: w=plain, b=bold, l=link."""
    out = []
    for m in UNIT_RE.finditer(text):
        s = m.group(0)
        lm = LINK_RE.fullmatch(s)
        if lm:
            out.append(("l", lm.group(1), lm.group(2)))
        elif s.startswith("**") and s.endswith("**") and len(s) > 4:
            out.append(("b", s[2:-2], None))
        else:
            out.append(("w", s, None))
    return out


def _clean_md_line(s: str) -> str:
    s = re.sub(r"!\[[^\]]*\]\([^)]*\)(\{[^}]*\})?", "", s)       # images
    s = re.sub(r"</?(u|span|sup|sub|div|br)[^>]*>", "", s)     # stray html
    s = re.sub(r"\{[#.][^}]*\}", "", s)                          # pandoc attrs
    s = re.sub(r"(?<!\*)\*(?!\*)([^*\n]+?)(?<!\*)\*(?!\*)", r"\1", s)  # *italic*
    s = re.sub(r"(?<![\w_])_([^_\n]+?)_(?![\w_])", r"\1", s)     # _italic_
    s = re.sub(r"\\([\\`*_{}\[\]()#+\-.!$<>|~'\"&])", r"\1", s)  # md escapes
    return s.strip()


# ---------------------------------------------------------------------------
# extract
# ---------------------------------------------------------------------------

def _to_markdown(src: Path) -> str:
    ext = src.suffix.lower()
    if ext == ".pdf":
        txt = subprocess.run(["pdftotext", str(src), "-"], capture_output=True, text=True, check=True).stdout
        paras = [re.sub(r"\s*\n\s*", " ", p).strip() for p in re.split(r"\n\s*\n", txt)]
        return "\n\n".join(p for p in paras if p)
    if ext == ".txt":
        return src.read_text(encoding="utf-8", errors="replace")
    fmt = {".docx": "docx", ".md": "markdown", ".markdown": "markdown",
           ".html": "html", ".htm": "html", ".odt": "odt", ".rtf": "rtf"}.get(ext)
    if fmt is None:
        sys.exit(f"ERROR: can't read {ext} files. Use .docx, .md, .txt, .html or .pdf.")
    if shutil.which("pandoc") is None:
        if ext in (".md", ".markdown"):
            return src.read_text(encoding="utf-8", errors="replace")
        if ext == ".docx":
            return _docx_to_markdown(src)
        sys.exit("ERROR: pandoc is not installed; it's needed to read this file type.")
    r = subprocess.run(["pandoc", "-f", fmt, "-t", "gfm", "--wrap=none", str(src)],
                       capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit(f"ERROR: pandoc failed: {r.stderr.strip()}")
    return _flatten_html_tables(r.stdout)


def _docx_to_markdown(src: Path) -> str:
    """Fallback .docx reader when pandoc is missing: headings, bullets, bold and links."""
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(str(src))
    out = []
    for p in d.paragraphs:
        parts = []
        for child in p._p.iterchildren():
            if child.tag == qn("w:r"):
                txt = "".join(t.text or "" for t in child.iter(qn("w:t")))
                bold = child.find(qn("w:rPr") + "/" + qn("w:b")) is not None
                parts.append(f"**{txt}**" if bold and txt.strip() else txt)
            elif child.tag == qn("w:hyperlink"):
                txt = "".join(t.text or "" for t in child.iter(qn("w:t")))
                rid = child.get(qn("r:id"))
                url = p.part.rels[rid].target_ref if rid in p.part.rels else ""
                parts.append(f"[{txt}]({url})" if url.startswith("http") else txt)
        text = "".join(parts).replace("****", "").strip()
        if not text:
            continue
        style = (p.style.name or "").lower()
        m = re.match(r"heading (\d)", style)
        if m or style == "title":
            out.append("#" * (int(m.group(1)) if m else 1) + " " + text.replace("**", ""))
        elif "list" in style:
            out.append("- " + text)
        else:
            out.append(text)
    for t in d.tables:
        out.append("\n".join("| " + " | ".join(c.text.strip().replace("|", "/") for c in r.cells) + " |"
                              + ("\n|" + "---|" * len(r.cells) if i == 0 else "")
                              for i, r in enumerate(t.rows)))
    return "\n\n".join(out)


def _html_to_md(fragment: str) -> str:
    r = subprocess.run(["pandoc", "-f", "html", "-t", "gfm", "--wrap=none"],
                       input=fragment, capture_output=True, text=True)
    return r.stdout.strip()


def _flatten_html_tables(md: str) -> str:
    """Pandoc writes tables it can't express in GFM (merged cells, several paragraphs
    in a cell) as raw HTML. A one-column table is usually a boxed note or header
    panel, so its cells become ordinary paragraphs; wider tables become pipe rows."""
    def repl(m):
        tbl = m.group(0)
        rows = re.findall(r"<tr\b.*?</tr>", tbl, flags=re.S)
        cells = [re.findall(r"<t[hd]\b[^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
        if all(len(c) <= 1 for c in cells):
            return "\n\n" + "\n\n".join(_html_to_md(c[0]) for c in cells if c) + "\n\n"
        lines = []
        for i, c in enumerate(cells):
            lines.append("| " + " | ".join(re.sub(r"\s+", " ", _html_to_md(x)).replace("|", "/") for x in c) + " |")
            if i == 0:
                lines.append("|" + "---|" * len(c))
        return "\n\n" + "\n".join(lines) + "\n\n"
    return re.sub(r"<table\b.*?</table>", repl, md, flags=re.S)


def parse_markdown(md: str) -> tuple[list[dict], list[str]]:
    """Markdown -> paragraphs [{n, style, text}] plus notes about what was simplified."""
    paras, notes, table = [], set(), []
    lines = md.replace("\r\n", "\n").split("\n")
    buf: list[str] = []

    def flush_para():
        if buf:
            text = _clean_md_line(" ".join(x.strip() for x in buf))
            if text:
                paras.append({"style": "p", "text": text})
            buf.clear()

    def flush_table():
        if table:
            rows = [[_clean_md_line(c) for c in r.strip().strip("|").split("|")]
                    for r in table if not re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", r)]
            paras.append({"style": "table", "rows": rows})
            notes.add("Tables are carried over unchanged (edits go in comments).")
            table.clear()

    for line in lines:
        if line.strip().startswith("|"):
            flush_para()
            table.append(line)
            continue
        flush_table()
        if not line.strip():
            flush_para()
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        li = re.match(r"^\s*(?:[-*+]|\d+[.)])\s+(.*)$", line)
        if h:
            flush_para()
            level = min(len(h.group(1)), 3)
            if len(h.group(1)) > 3:
                notes.add("Headings below H3 were treated as H3.")
            paras.append({"style": f"h{level}", "text": _clean_md_line(h.group(2))})
        elif li:
            flush_para()
            if re.match(r"^\s*\d+[.)]", line):
                notes.add("Numbered lists are shown as bullets.")
            paras.append({"style": "li", "text": _clean_md_line(li.group(1))})
        elif line.lstrip().startswith(">"):
            buf.append(line.lstrip()[1:])
        elif re.fullmatch(r"\s*([-*_]\s*){3,}", line):
            flush_para()
        else:
            buf.append(line)
    flush_para()
    flush_table()
    if re.search(r"!\[", md):
        notes.add("Images were left out; the review covers text only.")
    if re.search(r"(?<!\*)\*[^*\s][^*]*\*(?!\*)|(?<![\w_])_[^_\s][^_]*_", md):
        notes.add("Italics were dropped in the review copy.")
    for i, p in enumerate(paras, 1):
        p["n"] = i
    return paras, sorted(notes)


def cmd_extract(src: str, work: str) -> None:
    src_p, work_p = Path(src), Path(work)
    work_p.mkdir(parents=True, exist_ok=True)
    paras, notes = parse_markdown(_to_markdown(src_p))
    if not paras:
        sys.exit("ERROR: no text found in the input.")
    (work_p / "source.json").write_text(
        json.dumps({"source": src_p.name, "paragraphs": paras, "notes": notes}, indent=1, ensure_ascii=False),
        encoding="utf-8")
    words = sum(len(plain(p.get("text", "")).split()) for p in paras)
    print(f"# {src_p.name}: {len(paras)} paragraphs, about {words} words")
    for n in notes:
        print(f"# note: {n}")
    for p in paras:
        if p["style"] == "table":
            print(f"[{p['n']}] (table) " + " / ".join(" | ".join(r) for r in p["rows"]))
        else:
            print(f"[{p['n']}] ({p['style']}) {p['text']}")


# ---------------------------------------------------------------------------
# tracked
# ---------------------------------------------------------------------------

def _keep_tok(u):
    return ("kl", u[1], u[2]) if u[0] == "l" else ("kb", u[1]) if u[0] == "b" else ("k", u[1])


def _ins_tok(u):
    return ("il", u[1], u[2]) if u[0] == "l" else ("ib", u[1]) if u[0] == "b" else ("i", u[1])


def _merge(tokens):
    out = []
    for t in tokens:
        if out and t[0] in ("k", "i", "d", "kb", "ib") and out[-1][0] == t[0]:
            out[-1] = (t[0], out[-1][1] + t[1])
        else:
            out.append(t)
    return out


def diff_tokens(orig: str, rev: str) -> list[tuple]:
    """Word-level tracked tokens turning orig into rev. Tiny kept islands between two
    changes are folded into the change so the redline reads as phrases, not confetti."""
    a, b = units(orig), units(rev)
    ops = difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes()
    # Fold short equal runs (<= 2 units, e.g. " the") sandwiched between changes.
    merged = []
    for i, op in enumerate(ops):
        tag, i1, i2, j1, j2 = op
        if (tag == "equal" and 0 < i < len(ops) - 1 and (i2 - i1) <= 2
                and ops[i - 1][0] != "equal" and ops[i + 1][0] != "equal"):
            tag = "replace"
        if merged and tag != "equal" and merged[-1][0] != "equal":
            _, pi1, _, pj1, _ = merged[-1]
            merged[-1] = ("replace", pi1, i2, pj1, j2)
        else:
            merged.append((tag, i1, i2, j1, j2))
    toks = []
    for tag, i1, i2, j1, j2 in merged:
        if tag == "equal":
            toks += [_keep_tok(u) for u in a[i1:i2]]
            continue
        dels, ins = a[i1:i2], b[j1:j2]
        # Keep shared leading/trailing whitespace outside the change.
        lead = []
        while dels and ins and dels[0] == ins[0] and dels[0][0] == "w" and not dels[0][1].strip():
            lead.append(dels.pop(0)); ins.pop(0)
        trail = []
        while dels and ins and dels[-1] == ins[-1] and dels[-1][0] == "w" and not dels[-1][1].strip():
            trail.insert(0, dels.pop()); ins.pop()
        toks += [_keep_tok(u) for u in lead]
        toks += [("d", u[1]) for u in dels]
        toks += [_ins_tok(u) for u in ins]
        toks += [_keep_tok(u) for u in trail]
    return _merge(toks)


def _static_tokens(text):
    return _merge([_keep_tok(u) for u in units(text)])


def _ins_tokens(text):
    # Whole-paragraph insert: the builder wraps every run, so keep kinds are fine.
    return _merge([_keep_tok(u) for u in units(text)])


def build_blocks(source: dict, edits: dict):
    paras = {p["n"]: p for p in source["paragraphs"]}
    by_n, inserts = {}, {}
    for e in edits.get("edits", []):
        if "n" in e:
            if e["n"] not in paras:
                sys.exit(f"ERROR: edits.json refers to paragraph {e['n']}, which doesn't exist.")
            if e["n"] in by_n:
                sys.exit(f"ERROR: paragraph {e['n']} has two edits; combine them into one.")
            by_n[e["n"]] = e
        elif "after" in e:
            inserts.setdefault(e["after"], []).append(e)
        else:
            sys.exit(f"ERROR: each edit needs 'n' or 'after': {e}")
    comments = {int(k): tuple(v) for k, v in edits.get("comments", {}).items()}
    used = set()

    def cids(e):
        ids = [int(c) for c in e.get("c", [])]
        for c in ids:
            if c not in comments:
                sys.exit(f"ERROR: comment {c} is used but not defined in 'comments'.")
            used.add(c)
        return ids

    def with_comments(style, op, ids, toks):
        first = ids[0] if ids else None
        return (style, op, first, toks + [("c", c) for c in ids[1:]])

    blocks, expect_reject, expect_accept = [], [], []

    def emit_inserts(after):
        for e in inserts.get(after, []):
            style = e.get("style", "p")
            blocks.append(with_comments(style, "ins", cids(e), _ins_tokens(e["r"])))
            expect_accept.append(plain(e["r"]))

    emit_inserts(0)
    for n in sorted(paras):
        p = paras[n]
        e = by_n.get(n)
        if p["style"] == "table":
            blocks.append(("table", None, None, p["rows"]))
            if e:
                ids = cids(e)
                if ids:  # comment on a table: anchor it to an empty inserted note line
                    blocks.append(("p", None, ids[0], [("c", c) for c in ids[1:]]))
        elif e is None:
            blocks.append((p["style"], None, None, _static_tokens(p["text"])))
            expect_reject.append(plain(p["text"])); expect_accept.append(plain(p["text"]))
        else:
            ids = cids(e)
            orig, style = p["text"], e.get("style", p["style"])
            expect_reject.append(plain(orig))
            if e.get("delete"):
                blocks.append(with_comments(p["style"], "del", ids, _static_tokens(orig)))
            elif "r" not in e or e["r"] == orig and style == p["style"]:
                blocks.append(with_comments(p["style"], None, ids, _static_tokens(orig)))
                expect_accept.append(plain(orig))
            else:
                rev = e["r"]
                expect_accept.append(plain(rev))
                ratio = difflib.SequenceMatcher(None, units(orig), units(rev), autojunk=False).ratio()
                if style != p["style"] or ratio < 0.45:
                    # Heavy rewrite or style change: old paragraph out, new one in.
                    blocks.append((p["style"], "del", None, _static_tokens(orig)))
                    blocks.append(with_comments(style, "ins", ids, _ins_tokens(rev)))
                else:
                    blocks.append(with_comments(style, None, ids, diff_tokens(orig, rev)))
        emit_inserts(n)
    unused = set(comments) - used
    if unused:
        sys.exit(f"ERROR: comments {sorted(unused)} are defined but not attached to any edit.")
    return blocks, comments, expect_reject, expect_accept


def _simulate(blocks, mode):
    """Text a reader sees after Reject all (mode='reject') or Accept all ('accept')."""
    out = []
    for style, op, _cid, toks in blocks:
        if style == "table":
            continue
        if (mode == "reject" and op == "ins") or (mode == "accept" and op == "del"):
            continue
        keep = ("k", "kb", "kl", "d") if mode == "reject" else ("k", "kb", "kl", "i", "ib", "il")
        if op in ("ins", "del"):
            keep = ("k", "kb", "kl", "i", "ib", "il", "d")
        txt = "".join(t[1] for t in toks if t[0] in keep)
        if not txt.strip() and op is None and all(t[0] == "c" for t in toks):
            continue  # comment-only anchor line for a table
        out.append(re.sub(r"\s+", " ", txt).strip())
    return out


def slim_docx(path: Path) -> None:
    """Drop parts python-docx's default template carries but never uses (a thumbnail,
    a duplicate style sheet, sample custom XML) and prune styles.xml to the styles the
    document references. Cuts ~40KB to ~10KB, which matters because a Drive upload
    through the connector passes the file as base64 text."""
    drop = {"word/stylesWithEffects.xml", "docProps/thumbnail.jpeg", "word/webSettings.xml",
            "customXml/item1.xml", "customXml/_rels/item1.xml.rels", "customXml/itemProps1.xml"}
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist() if n not in drop}
    ct = files["[Content_Types].xml"].decode()
    for part in drop:
        ct = re.sub(r'<Override PartName="/' + re.escape(part) + r'"[^>]*/>', "", ct)
    files["[Content_Types].xml"] = ct.encode()
    for relname in ("_rels/.rels", "word/_rels/document.xml.rels"):
        rels = files[relname].decode()
        rels = re.sub(r'<Relationship [^>]*Target="[^"]*(thumbnail\.jpeg|stylesWithEffects\.xml|'
                      r'webSettings\.xml|customXml/item1\.xml)"[^>]*/>', "", rels)
        files[relname] = rels.encode()

    styles = files["word/styles.xml"].decode()
    used = set(re.findall(r'w:(?:pStyle|rStyle|tblStyle) w:val="([^"]+)"',
                          files["word/document.xml"].decode() + files.get("word/comments.xml", b"").decode()))
    blocks = re.findall(r"<w:style\b.*?</w:style>", styles, flags=re.S)
    by_id = {}
    for b in blocks:
        m = re.search(r'w:styleId="([^"]+)"', b)
        if m:
            by_id[m.group(1)] = b
    keep = set(used) | {i for i, b in by_id.items() if 'w:default="1"' in b}
    changed = True
    while changed:  # follow basedOn / link / next chains
        changed = False
        for sid in list(keep):
            for ref in re.findall(r'<w:(?:basedOn|link|next) w:val="([^"]+)"', by_id.get(sid, "")):
                if ref not in keep:
                    keep.add(ref); changed = True
    for sid, b in by_id.items():
        if sid not in keep:
            styles = styles.replace(b, "")
    styles = re.sub(r"<w:latentStyles\b.*?</w:latentStyles>", "", styles, flags=re.S)
    files["word/styles.xml"] = styles.encode()
    tmp = path.with_suffix(".slim.tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for n, data in files.items():
            z.writestr(n, data)
    tmp.replace(path)


def verify_docx(path: Path, n_comments: int) -> list[str]:
    problems = []
    with zipfile.ZipFile(path) as z:
        doc = z.read("word/document.xml").decode()
        cxml = z.read("word/comments.xml").decode() if "word/comments.xml" in z.namelist() else ""
    if "‖" in doc:
        problems.append("a comment marker was left in the text (a comment failed to anchor)")
    refs = len(set(re.findall(r'<w:commentReference w:id="(\d+)"', doc)))
    if refs != n_comments:
        problems.append(f"{n_comments} comments defined but {refs} anchored in the document")
    if cxml.count("<w:comment ") != n_comments:
        problems.append("comments.xml doesn't hold every comment")
    return problems


def cmd_tracked(work: str, out: str) -> None:
    import build_review_tracked as brt

    work_p, out_p = Path(work), Path(out)
    source = json.loads((work_p / "source.json").read_text(encoding="utf-8"))
    edits = json.loads((work_p / "edits.json").read_text(encoding="utf-8"))
    blocks, comments, exp_rej, exp_acc = build_blocks(source, edits)

    problems = []
    if _simulate(blocks, "reject") != exp_rej:
        problems.append("Reject all would not restore the original text")
    if _simulate(blocks, "accept") != exp_acc:
        problems.append("Accept all would not produce the revised text")

    brt.BLOCKS, brt.COMMENTS = blocks, comments
    brt.TITLE = edits.get("title") or source.get("source", "Draft")
    brt.AUTHOR = edits.get("author", "CopyChief (Editor)")
    out_p.parent.mkdir(parents=True, exist_ok=True)
    brt.build_docx(out_p, header=False)
    brt.build_html(out_p.with_suffix(".html"))
    slim_docx(out_p)
    problems += verify_docx(out_p, len(comments))

    changed = sum(1 for b in blocks if b[1] or any(t[0] in ("i", "d", "ib", "il") for t in b[3] if b[0] != "table"))
    print(f"Wrote: {out_p} ({out_p.stat().st_size // 1024} KB)")
    print(f"Wrote: {out_p.with_suffix('.html')}")
    print(f"{changed} paragraphs changed, {len(comments)} comments")
    if problems:
        print("VERIFY FAILED:\n  - " + "\n  - ".join(problems))
        sys.exit(1)
    print("VERIFY OK: reject-all restores the original; accept-all gives clean revised copy; all comments anchored.")


# ---------------------------------------------------------------------------
# doc / b64
# ---------------------------------------------------------------------------

def cmd_doc(blocks_json: str, out: str) -> None:
    import docx_lib

    spec = json.loads(Path(blocks_json).read_text(encoding="utf-8"))
    blocks = [tuple(b) for b in spec["blocks"]]
    comments = {int(k): tuple(v) for k, v in spec.get("comments", {}).items()}
    title = spec.get("title")
    out_p = Path(out)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    doc = docx_lib.render_docx(blocks, comments=comments or None, title=title)
    doc.save(out_p)
    slim_docx(out_p)
    out_p.with_suffix(".md").write_text(docx_lib.render_md(blocks, title=title), encoding="utf-8")
    missing = [c for c in comments if "{{c%d}}" % c not in json.dumps(spec["blocks"])]
    print(f"Wrote: {out_p} ({out_p.stat().st_size // 1024} KB)")
    print(f"Wrote: {out_p.with_suffix('.md')}")
    if missing:
        print(f"WARNING: comments {missing} are defined but never anchored with {{{{cN}}}}.")
        sys.exit(1)


def cmd_b64(path: str) -> None:
    sys.stdout.write(base64.b64encode(Path(path).read_bytes()).decode())


def main() -> None:
    cmds = {"extract": (cmd_extract, 2), "tracked": (cmd_tracked, 2), "doc": (cmd_doc, 2), "b64": (cmd_b64, 1)}
    if len(sys.argv) < 2 or sys.argv[1] not in cmds or len(sys.argv) - 2 != cmds[sys.argv[1]][1]:
        sys.exit(__doc__)
    fn, _ = cmds[sys.argv[1]]
    fn(*sys.argv[2:])


if __name__ == "__main__":
    main()
