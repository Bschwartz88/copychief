"""Tracked-changes review builder for CopyChief (client folders and the projects workspace).

Builds an editorial review .docx (native Word tracked changes + anchored sidebar
comments) and an HTML redline for any project. Client-neutral: the article content
and editorial notes live in a per-project data file, not in this engine.

Why tracked changes: a .docx with native Word tracked changes (w:ins / w:del) plus
comments, when uploaded to Google Drive and opened as a Google Doc, becomes a doc
with real *suggestions* (accept/reject one by one) and working margin comments.
Highlights alone cannot be accepted/rejected; the Docs/Drive API cannot create
suggestions, so a tracked-changes .docx is the only reliable path.

Design (do not change without reading the README):
  - Base text = the ORIGINAL draft. Each change is a w:del of the original span plus
    a w:ins of the replacement, so Accept keeps the edit and Reject restores the
    original wording. Whole-paragraph rewrites also mark the paragraph mark so
    accept/reject merges cleanly with no blank lines.
  - Rationale lives ONLY in anchored Word comments (not inline superscripts), so
    "Accept all" in Google Docs leaves clean final copy with no review artifacts.

Usage:
  python tools/build_review_tracked.py <slug>

Reads:   the slug's data file (see _paths.py) -- a TRACKED dict, or BLOCKS/COMMENTS
Writes:  <reviews folder>/<slug>-edits.docx   (upload to Google Drive)
         <reviews folder>/<slug>-edits.html   (visual redline, reference only)

Google Drive workflow:
  1. Upload the .docx to Drive.
  2. Right-click -> Open with -> Google Docs (or upload with "Convert uploads" on).
  3. Tracked changes appear as suggestions; comments appear in the margin.
  4. Accept/reject each suggestion; resolve comments as you go.

The per-project data file's content model:
  BLOCKS = list of tuples, each either:
    (style, op, comment_id, tokens)      for text paragraphs/headings/list items
    ("table", None, None, rows)          for a table (rows = list of cell-lists)
  style:  "h1" "h2" "h3" "p" "li"
  op:     None (in-place token-level edits) | "ins" (whole para inserted) | "del" (whole para deleted)
  comment_id: an int key into COMMENTS, or None
  Token kinds inside a paragraph:
    ("k", text)        keep                ("kb", text)  keep bold
    ("i", text)        insert              ("ib", text)  insert bold
    ("d", text)        delete
    ("kl", text, url)  keep link           ("il", text, url) insert link
    ("c", cid)         point comment (zero-width anchor, no text change)
  COMMENTS = dict: cid -> (short label, full reasoning)
  TITLE  (optional str)  document/redline heading. Defaults to the slug.
  AUTHOR (optional str)  tracked-change / comment author. Defaults to "CopyChief (Editor)".
"""

from __future__ import annotations

import html as htmllib
import importlib.util
import itertools
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DATE = "2026-01-01T00:00:00Z"  # fixed revision date keeps output deterministic
_rev = itertools.count(1000)   # tracked-change revision ids (distinct from comment ids)

# Populated from the per-project data file at runtime.
BLOCKS: list[tuple] = []
COMMENTS: dict[int, tuple[str, str]] = {}
AUTHOR = "CopyChief (Editor)"
TITLE = ""


# Sentinel chars (U+2016) for comment-range post-processing.
def smark(cid: int) -> str: return f"‖C{cid}S‖"
def emark(cid: int) -> str: return f"‖C{cid}E‖"
def pmark(cid: int) -> str: return f"‖C{cid}P‖"


# ---------------------------------------------------------------------------
# Load the per-project data file
# ---------------------------------------------------------------------------

def load_project(slug: str) -> Path:
    """Load the slug's data file (see _paths.py) and pull BLOCKS / COMMENTS / TITLE / AUTHOR.

    The data file may define those names at module level, or a TRACKED dict with
    keys "blocks", "comments", and optional "title" / "author" (so one data file
    can hold PREP, REVIEW and TRACKED side by side).
    """
    global BLOCKS, COMMENTS, AUTHOR, TITLE
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from _paths import load_data, resolve
    mod = load_data(slug)
    src = getattr(mod, "TRACKED", None)
    if src is not None:
        BLOCKS, COMMENTS = src["blocks"], src["comments"]
        AUTHOR = src.get("author", "CopyChief (Editor)")
        TITLE = src.get("title", slug)
    else:
        if not hasattr(mod, "BLOCKS"):
            sys.exit(f"ERROR: {resolve(slug)['data']} defines neither TRACKED nor BLOCKS.")
        BLOCKS, COMMENTS = mod.BLOCKS, mod.COMMENTS
        AUTHOR = getattr(mod, "AUTHOR", "CopyChief (Editor)")
        TITLE = getattr(mod, "TITLE", slug)
    return resolve(slug)["reviews"]


# ---------------------------------------------------------------------------
# DOCX: tracked-changes helpers (direct OOXML)
# ---------------------------------------------------------------------------

def _wrap(el, tag: str) -> None:
    """Wrap an existing run/hyperlink element in a w:ins or w:del revision element."""
    w = OxmlElement(tag)
    w.set(qn("w:id"), str(next(_rev)))
    w.set(qn("w:author"), AUTHOR)
    w.set(qn("w:date"), DATE)
    el.addprevious(w)
    w.append(el)


def _add_hyperlink(paragraph, url: str, text: str):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), "0563C1")
    r_pr.append(color_el)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _keep_run(p, text, bold=False):
    r = p.add_run(text)
    if bold:
        r.bold = True
    return r


def _ins_run(p, text, bold=False):
    r = p.add_run(text)
    if bold:
        r.bold = True
    _wrap(r._r, "w:ins")


def _del_run(p, text, bold=False):
    r = p.add_run(text)
    if bold:
        r.bold = True
    t = r._r.find(qn("w:t"))
    if t is not None:
        t.tag = qn("w:delText")
    _wrap(r._r, "w:del")


def _ins_hyperlink(p, url, text):
    hl = _add_hyperlink(p, url, text)
    _wrap(hl, "w:ins")


def _sentinel(p, marker):
    p.add_run(marker)


def _mark_para(p, tag: str) -> None:
    """Mark the paragraph mark as inserted (w:ins) or deleted (w:del)."""
    pPr = p._p.get_or_add_pPr()
    rPr = pPr.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        pPr.append(rPr)
    el = OxmlElement(tag)
    el.set(qn("w:id"), str(next(_rev)))
    el.set(qn("w:author"), AUTHOR)
    el.set(qn("w:date"), DATE)
    rPr.append(el)


_STYLE = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "p": "Normal", "li": "List Bullet"}


def _render_paragraph(doc, style, op, cid, tokens):
    p = doc.add_paragraph(style=_STYLE.get(style, "Normal"))
    if cid is not None:
        _sentinel(p, smark(cid))
    for tok in tokens:
        kind = tok[0]
        if kind == "c":
            _sentinel(p, pmark(tok[1]))
            continue
        if op == "ins":
            if kind in ("kl", "il"):
                _ins_hyperlink(p, tok[2], tok[1])
            else:
                _ins_run(p, tok[1], bold=kind in ("kb", "ib"))
        elif op == "del":
            # Links inside a deleted paragraph render as plain struck text.
            _del_run(p, tok[1], bold=kind == "kb")
        else:
            if kind == "k":
                _keep_run(p, tok[1])
            elif kind == "kb":
                _keep_run(p, tok[1], bold=True)
            elif kind == "i":
                _ins_run(p, tok[1])
            elif kind == "ib":
                _ins_run(p, tok[1], bold=True)
            elif kind == "d":
                _del_run(p, tok[1])
            elif kind == "kl":
                _add_hyperlink(p, tok[2], tok[1])
            elif kind == "il":
                _ins_hyperlink(p, tok[2], tok[1])
    if cid is not None:
        _sentinel(p, emark(cid))
    if op == "ins":
        _mark_para(p, "w:ins")
    elif op == "del":
        _mark_para(p, "w:del")


def _render_table(doc, rows):
    ncol = max(len(r) for r in rows)
    tbl = doc.add_table(rows=0, cols=ncol)
    tbl.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            run = cells[ci].paragraphs[0].add_run(txt)
            run.font.size = Pt(10)
            if ri == 0:
                run.bold = True


def build_docx(docx_path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    t = doc.add_paragraph()
    tr = t.add_run(f"Editorial Review (tracked changes) — {TITLE}")
    tr.bold = True
    tr.font.size = Pt(15)
    s = doc.add_paragraph()
    sr = s.add_run(
        f"Reviewed {datetime.now():%Y-%m-%d}. Edits are Word tracked changes; rationale is in the "
        f"anchored comments. Upload to Google Drive and open as a Google Doc to accept/reject each "
        f"suggestion with comments in the margin."
    )
    sr.italic = True
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph()

    for style, op, cid, payload in BLOCKS:
        if style == "table":
            _render_table(doc, payload)
        else:
            _render_paragraph(doc, style, op, cid, payload)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = docx_path.parent / "_review_tmp.docx"
    doc.save(tmp)
    _postprocess_comments(tmp, docx_path)
    tmp.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# Post-process: inject native Word comments at the sentinel ranges
# ---------------------------------------------------------------------------

def _xml_escape(text: str) -> str:
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


def _replace_sentinels(doc_xml: str) -> str:
    def repl(sentinel, replacement, xml):
        pat = re.compile(
            r"<w:r\b(?:(?!</w:r>).)*?<w:t[^>]*>" + re.escape(sentinel) + r"</w:t>(?:(?!</w:r>).)*?</w:r>",
            flags=re.DOTALL,
        )
        return pat.sub(replacement, xml)

    for cid in COMMENTS:
        doc_xml = repl(smark(cid), f'<w:commentRangeStart w:id="{cid}"/>', doc_xml)
        doc_xml = repl(
            emark(cid),
            f'<w:commentRangeEnd w:id="{cid}"/>'
            f'<w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>'
            f'<w:commentReference w:id="{cid}"/></w:r>',
            doc_xml,
        )
        doc_xml = repl(
            pmark(cid),
            f'<w:commentRangeStart w:id="{cid}"/><w:commentRangeEnd w:id="{cid}"/>'
            f'<w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>'
            f'<w:commentReference w:id="{cid}"/></w:r>',
            doc_xml,
        )
    return doc_xml


def _comments_xml() -> str:
    parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">',
    ]
    today = datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ")
    for cid, (label, reason) in COMMENTS.items():
        body = f"{label}. {reason}"
        parts.append(
            f'<w:comment w:id="{cid}" w:author="{_xml_escape(AUTHOR)}" w:date="{today}" w:initials="CC">'
            f'<w:p><w:r><w:t xml:space="preserve">{_xml_escape(body)}</w:t></w:r></w:p></w:comment>'
        )
    parts.append("</w:comments>")
    return "".join(parts)


def _postprocess_comments(src: Path, dst: Path) -> None:
    with zipfile.ZipFile(src, "r") as zin:
        files = {n: zin.read(n) for n in zin.namelist()}

    files["word/document.xml"] = _replace_sentinels(
        files["word/document.xml"].decode("utf-8")).encode("utf-8")
    files["word/comments.xml"] = _comments_xml().encode("utf-8")

    ct = files["[Content_Types].xml"].decode("utf-8")
    override = ('<Override PartName="/word/comments.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>')
    if override not in ct:
        ct = ct.replace("</Types>", override + "</Types>")
    files["[Content_Types].xml"] = ct.encode("utf-8")

    rels = files["word/_rels/document.xml.rels"].decode("utf-8")
    if "comments.xml" not in rels:
        ids = [int(i) for i in re.findall(r'Id="rId(\d+)"', rels)]
        nid = (max(ids) + 1) if ids else 1
        rel = (f'<Relationship Id="rId{nid}" '
               f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" '
               f'Target="comments.xml"/>')
        rels = rels.replace("</Relationships>", rel + "</Relationships>")
    files["word/_rels/document.xml.rels"] = rels.encode("utf-8")

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in files.items():
            zout.writestr(n, data)


# ---------------------------------------------------------------------------
# HTML redline (reference only — not for Drive accept/reject)
# ---------------------------------------------------------------------------

HTML_CSS = """
body{font-family:-apple-system,"Segoe UI",Calibri,sans-serif;max-width:840px;margin:2rem auto;
     padding:0 1.5rem;line-height:1.55;color:#1a1a1a;}
h1{font-size:1.6rem;} h2{font-size:1.3rem;margin-top:1.8rem;} h3{font-size:1.05rem;}
ins{background:#e6ffed;text-decoration:underline;text-decoration-color:#22863a;}
del{background:#ffeef0;text-decoration:line-through;color:#b31d28;}
a{color:#0563c1;}
table{border-collapse:collapse;margin:1rem 0;width:100%;}
th,td{border:1px solid #b8c2cc;padding:6px 9px;text-align:left;vertical-align:top;font-size:.92rem;}
th{background:#eef3f8;}
.cmt{display:block;margin:.25rem 0 .9rem;padding:.5rem .75rem;background:#fff8e1;border-left:4px solid #f0c000;
     font-size:.86rem;color:#5b4a00;border-radius:3px;}
.cmt b{color:#3a3000;}
.meta{font-size:.85rem;color:#555;font-style:italic;}
.banner{background:#eef3f8;border:1px solid #b8c2cc;border-radius:4px;padding:.75rem 1rem;font-size:.9rem;}
"""


def _html_tokens(op, tokens):
    out = []
    for tok in tokens:
        kind = tok[0]
        if kind == "c":
            continue
        if op == "ins":
            if kind in ("kl", "il"):
                out.append(f'<ins><a href="{htmllib.escape(tok[2])}">{htmllib.escape(tok[1])}</a></ins>')
            else:
                txt = htmllib.escape(tok[1])
                out.append(f"<ins>{'<b>'+txt+'</b>' if kind in ('kb','ib') else txt}</ins>")
        elif op == "del":
            out.append(f"<del>{htmllib.escape(tok[1])}</del>")
        else:
            if kind == "k":
                out.append(htmllib.escape(tok[1]))
            elif kind == "kb":
                out.append(f"<b>{htmllib.escape(tok[1])}</b>")
            elif kind == "i":
                out.append(f"<ins>{htmllib.escape(tok[1])}</ins>")
            elif kind == "ib":
                out.append(f"<ins><b>{htmllib.escape(tok[1])}</b></ins>")
            elif kind == "d":
                out.append(f"<del>{htmllib.escape(tok[1])}</del>")
            elif kind == "kl":
                out.append(f'<a href="{htmllib.escape(tok[2])}">{htmllib.escape(tok[1])}</a>')
            elif kind == "il":
                out.append(f'<ins><a href="{htmllib.escape(tok[2])}">{htmllib.escape(tok[1])}</a></ins>')
    return "".join(out)


def _comment_html(cid):
    if cid is None:
        return ""
    label, reason = COMMENTS[cid]
    return f'<span class="cmt"><b>Comment — {htmllib.escape(label)}:</b> {htmllib.escape(reason)}</span>'


def build_html(html_path: Path) -> None:
    parts, in_list = [], False
    for style, op, cid, payload in BLOCKS:
        if style != "li" and in_list:
            parts.append("</ul>")
            in_list = False
        if style == "table":
            rows = payload
            parts.append("<table>")
            for ri, row in enumerate(rows):
                tag = "th" if ri == 0 else "td"
                parts.append("<tr>" + "".join(f"<{tag}>{htmllib.escape(c)}</{tag}>" for c in row) + "</tr>")
            parts.append("</table>")
            continue
        inner = _html_tokens(op, payload)
        point_cids = [t[1] for t in payload if t and t[0] == "c"]
        cmt = _comment_html(cid) + "".join(_comment_html(pc) for pc in point_cids)
        if style == "h1":
            parts.append(f"<h1>{inner}</h1>{cmt}")
        elif style == "h2":
            parts.append(f"<h2>{inner}</h2>{cmt}")
        elif style == "h3":
            parts.append(f"<h3>{inner}</h3>{cmt}")
        elif style == "li":
            if not in_list:
                parts.append("<ul>")
                in_list = True
            parts.append(f"<li>{inner}{cmt}</li>")
        else:
            parts.append(f"<p>{inner}</p>{cmt}")
    if in_list:
        parts.append("</ul>")

    today = datetime.now().strftime("%Y-%m-%d")
    title = htmllib.escape(TITLE)
    html = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>Editorial redline — {title}</title><style>{HTML_CSS}</style></head><body>
<h1 style="color:#1f4e79;">Editorial redline — {title}</h1>
<p class="meta">Reviewed {today}.</p>
<p class="banner"><b>Reference only.</b> This HTML shows the redline visually
(<ins>green = insert</ins>, <del>red = delete</del>) with rationale in the yellow comment boxes.
For accept/reject in Google Drive, use the <code>.docx</code> in this folder: upload it and open as a
Google Doc, where the tracked changes become suggestions and the comments appear in the margin.</p>
{''.join(parts)}
</body></html>"""
    html_path.write_text(html, encoding="utf-8")


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("Usage: python tools/build_review_tracked.py <slug>")
    slug = sys.argv[1].strip().strip("/\\")
    reviews = load_project(slug)
    reviews.mkdir(parents=True, exist_ok=True)
    docx_path = reviews / f"{slug}-edits.docx"
    html_path = reviews / f"{slug}-edits.html"
    build_docx(docx_path)
    build_html(html_path)
    print(f"Wrote: {docx_path}")
    print(f"Wrote: {html_path}")


if __name__ == "__main__":
    main()
